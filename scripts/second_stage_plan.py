from __future__ import annotations

import re
import unicodedata
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from common import PROCESSED_FIELD_MAP, stringify, to_int, write_csv_rows, write_json


HEADER_ALIASES = {
    "school_code": {"院校代码", "学校代码", "院校代号"},
    "school_name": {"院校名称", "学校名称"},
    "major_code": {"专业代码", "专业代号"},
    "major_name": {"专业名称", "专业(类)名称", "专业（类）名称", "招生专业"},
    "plan_count": {"计划数", "剩余计划数", "招生计划数", "招生计划"},
    "tuition": {"学费", "学费参考", "收费标准"},
    "duration": {"学制", "修业年限"},
    "campus": {"校区", "办学地点", "就读地点"},
    "city": {"城市", "所在地", "办学地"},
    "school_province": {"省", "省份", "院校省份", "学校省份"},
    "subject_requirement": {"选考科目要求", "选考科目范围", "选科要求", "选考要求"},
    "foreign_language_requirement": {"外语语种要求", "外语语种", "语种要求"},
    "single_subject_requirement": {"单科成绩要求", "单科要求"},
    "gender_requirement": {"性别要求", "性别限制"},
    "medical_requirement": {"体检要求", "身体条件要求"},
    "remarks": {"专业备注", "专业简注", "备注", "说明"},
    "degree_level": {"层次", "专业层次", "学历层次", "培养层次", "本专科"},
    "school_type": {"办学性质", "院校性质", "学校性质"},
    "application_stage": {"填报段次", "段次", "批次"},
}

REQUIRED_IMPORT_FIELDS = {
    "school_code",
    "school_name",
    "major_code",
    "major_name",
    "plan_count",
    "subject_requirement",
    "degree_level",
}

OFFICIAL_PLAN_PORTAL = "https://www.zjzs.net/col/col4/"


def normalize_cell(value: object) -> str:
    if value is None:
        return ""
    text = unicodedata.normalize("NFKC", str(value)).strip()
    return re.sub(r"\s+", "", text)


def normalize_level(value: object) -> str:
    text = normalize_cell(value)
    if "职业本科" in text:
        return "职业本科"
    if "本科" in text:
        return "本科"
    if "专科" in text or "高职" in text:
        return "专科"
    return ""


def normalize_stage(value: object) -> str:
    text = normalize_cell(value).replace("普通类", "")
    if "二段" in text:
        return "二段"
    if "一段" in text:
        return "一段"
    return ""


def header_field(value: object) -> str:
    normalized = normalize_cell(value)
    for field, aliases in HEADER_ALIASES.items():
        if normalized in {normalize_cell(alias) for alias in aliases}:
            return field
    return ""


def locate_header(workbook: object) -> tuple[object, int, dict[str, int]]:
    best: tuple[object, int, dict[str, int]] | None = None
    for sheet in workbook.worksheets:
        for row_index, row in enumerate(
            sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 40), values_only=True),
            start=1,
        ):
            mapping: dict[str, int] = {}
            for column_index, value in enumerate(row):
                field = header_field(value)
                if field and field not in mapping:
                    mapping[field] = column_index
            if best is None or len(mapping) > len(best[2]):
                best = (sheet, row_index, mapping)
            if REQUIRED_IMPORT_FIELDS <= set(mapping):
                return sheet, row_index, mapping
    if best is None:
        raise ValueError("工作簿中没有可读取的表格")
    missing = sorted(REQUIRED_IMPORT_FIELDS - set(best[2]))
    raise ValueError(
        "无法识别个性化二段计划表头，缺少："
        + "、".join(missing)
        + "。不得根据段次推断学历层次，请确认文件包含“层次/专业层次”列。"
    )


def cell(row: tuple[object, ...], mapping: dict[str, int], field: str) -> str:
    index = mapping.get(field)
    if index is None or index >= len(row):
        return ""
    return stringify(row[index]).strip()


def infer_language_requirement(remarks: str, explicit: str) -> str:
    if explicit:
        return explicit
    if "只招英语" in remarks or "限英语" in remarks:
        return "英语"
    return ""


def import_second_stage_plan(
    input_path: Path,
    output_path: Path,
    year: int = 2026,
) -> tuple[list[dict[str, object]], dict[str, object]]:
    if not input_path.exists():
        raise FileNotFoundError(
            f"缺少2026个性化二段计划文件：{input_path}。"
            f"请登录浙江省教育考试院“计划查询/平行志愿计划下载”后取得 XLSM/XLSX：{OFFICIAL_PLAN_PORTAL}"
        )
    if input_path.suffix.lower() not in {".xlsm", ".xlsx"}:
        raise ValueError("个性化二段计划必须是浙江考试院下载的 .xlsm 或 .xlsx 文件")
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("缺少 openpyxl，无法读取 XLSM/XLSX") from exc

    workbook = load_workbook(
        input_path,
        read_only=True,
        data_only=True,
        keep_vba=input_path.suffix.lower() == ".xlsm",
    )
    sheet, header_row, mapping = locate_header(workbook)
    rows: list[dict[str, object]] = []
    updated_at = datetime.now(timezone.utc).isoformat(timespec="seconds")
    seen: set[tuple[str, str, str]] = set()

    for values in sheet.iter_rows(min_row=header_row + 1, values_only=True):
        school_code = cell(values, mapping, "school_code")
        school_name = cell(values, mapping, "school_name")
        major_code = cell(values, mapping, "major_code")
        major_name = cell(values, mapping, "major_name")
        if not any([school_code, school_name, major_code, major_name]):
            continue
        if not all([school_code, school_name, major_code, major_name]):
            continue
        degree_level = normalize_level(cell(values, mapping, "degree_level"))
        if not degree_level:
            raise ValueError(f"{school_name} {major_name} 的学历层次无法识别，已停止导入")
        stage = normalize_stage(cell(values, mapping, "application_stage")) or "二段"
        if stage != "二段":
            raise ValueError(f"{school_name} {major_name} 的段次为{stage}，不是二段个性化剩余计划")
        plan_count = to_int(cell(values, mapping, "plan_count"))
        if plan_count is None or plan_count < 1:
            continue
        remarks = cell(values, mapping, "remarks")
        school_type = cell(values, mapping, "school_type")
        key = (school_code, major_code, cell(values, mapping, "campus"))
        if key in seen:
            continue
        seen.add(key)
        rows.append(
            {
                "year": year,
                "province": "浙江",
                "school_province": cell(values, mapping, "school_province"),
                "batch": "普通类二段",
                "application_stage": "二段",
                "school_code": school_code,
                "school_name": school_name,
                "major_code": major_code,
                "major_name": major_name,
                "major_group": "",
                "plan_count": plan_count,
                "tuition": to_int(cell(values, mapping, "tuition"), 0) or "",
                "duration": cell(values, mapping, "duration"),
                "campus": cell(values, mapping, "campus"),
                "city": cell(values, mapping, "city"),
                "subject_requirement": cell(values, mapping, "subject_requirement") or "不限",
                "foreign_language_requirement": infer_language_requirement(
                    remarks,
                    cell(values, mapping, "foreign_language_requirement"),
                ),
                "single_subject_requirement": cell(values, mapping, "single_subject_requirement"),
                "gender_requirement": cell(values, mapping, "gender_requirement"),
                "medical_requirement": cell(values, mapping, "medical_requirement"),
                "remarks": remarks,
                "is_public": "true" if "公办" in school_type else "",
                "is_private": "true" if "民办" in school_type else "",
                "is_sino_foreign": "true" if "中外合作" in (major_name + remarks + school_type) else "false",
                "degree_level": degree_level,
                "plan_status": "2026二段当前可填报",
                "source_access": "personalized-login-download",
                "source_file": input_path.name,
                "source_url": OFFICIAL_PLAN_PORTAL,
                "updated_at": updated_at,
            }
        )

    workbook.close()
    if not rows:
        raise ValueError("个性化二段计划文件中没有解析到可填报专业")
    write_csv_rows(output_path, rows, PROCESSED_FIELD_MAP["admissions_plan.csv"])
    level_counts = Counter(str(row["degree_level"]) for row in rows)
    summary = {
        "input_file": str(input_path),
        "output_file": str(output_path),
        "application_stage": "二段",
        "current_plan_verified": True,
        "program_count": len(rows),
        "seat_count": sum(int(row["plan_count"]) for row in rows),
        "undergraduate_program_count": level_counts["本科"] + level_counts["职业本科"],
        "associate_program_count": level_counts["专科"],
        "imported_at": updated_at,
    }
    write_json(output_path.parent / "second_stage_plan_import_summary.json", summary)
    return rows, summary
