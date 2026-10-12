from __future__ import annotations

import argparse
import json
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

from common import ROOT


def read_sheet(path: Path) -> tuple[str, list[dict[str, object]]]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = workbook.worksheets[0]
    values = list(sheet.iter_rows(values_only=True))
    headers = [str(value).strip() if value is not None else "" for value in values[0]]
    rows: list[dict[str, object]] = []
    for values_row in values[1:]:
        row = {
            header: value
            for header, value in zip(headers, values_row)
            if header
        }
        if any(value not in {None, ""} for value in row.values()):
            rows.append(row)
    workbook.close()
    return sheet.title, rows


def main() -> int:
    parser = argparse.ArgumentParser(description="提取两张候选工作簿供规则核验。")
    parser.add_argument(
        "--other-ai",
        type=Path,
        default=ROOT / "data" / "raw" / "other_ai_candidates.xlsx",
    )
    parser.add_argument(
        "--selected",
        type=Path,
        default=ROOT / "data" / "raw" / "selected_candidates.xlsx",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=ROOT / "outputs" / "input_review",
    )
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)

    other_sheet, other_rows = read_sheet(args.other_ai)
    selected_sheet, selected_rows = read_sheet(args.selected)
    (args.output_dir / "other_ai_rows.json").write_text(
        json.dumps(other_rows, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    (args.output_dir / "selected_rows.json").write_text(
        json.dumps(selected_rows, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    summary = {
        "other_ai": {
            "sheet": other_sheet,
            "rows": len(other_rows),
            "tiers": dict(Counter(str(row.get("梯度", "")) for row in other_rows)),
        },
        "selected": {
            "sheet": selected_sheet,
            "rows": len(selected_rows),
            "tiers": dict(Counter(str(row.get("层级", "")) for row in selected_rows)),
            "status": dict(
                Counter(str(row.get("填报状态", "")) for row in selected_rows)
            ),
        },
    }
    print(json.dumps(summary, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
